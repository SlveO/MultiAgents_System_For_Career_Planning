from __future__ import annotations

from pathlib import Path
from typing import List, Optional

try:
    from project.core.schemas import EvidenceItem, PerceptionResult
except ImportError:
    from core.schemas import EvidenceItem, PerceptionResult

from .base import _extract_facts_fallback, _safe_confidence


class DocumentPerceptionAgent:
    @staticmethod
    def _read_txt(path: Path) -> str:
        return path.read_text(encoding="utf-8", errors="ignore")

    @staticmethod
    def _read_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
        except Exception:
            return "pypdf not installed, cannot parse PDF."
        reader = PdfReader(str(path))
        chunks = []
        for page in reader.pages[:10]:
            chunks.append(page.extract_text() or "")
        return "\n".join(chunks)

    @staticmethod
    def _read_docx(path: Path) -> str:
        try:
            import docx
        except Exception:
            return "python-docx not installed, cannot parse DOCX."
        doc = docx.Document(str(path))
        text = "\n".join([p.text for p in doc.paragraphs]).strip()
        return text if text else f"[{path.name}] DOCX has no extractable text."

    @staticmethod
    def _read_xlsx(path: Path) -> str:
        try:
            import openpyxl
        except Exception:
            return "openpyxl not installed, cannot parse XLSX."
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        lines: List[str] = []
        for ws in wb.worksheets[:3]:
            lines.append(f"[sheet] {ws.title}")
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                values = [str(x) for x in row if x is not None and str(x).strip()]
                if values:
                    lines.append(" | ".join(values[:12]))
                row_count += 1
                if row_count >= 20:
                    break
        wb.close()
        return "\n".join(lines).strip() if lines else f"[{path.name}] XLSX has no extractable content."

    def perceive(self, document_path: str) -> PerceptionResult:
        path = Path(document_path)
        if not path.exists():
            return PerceptionResult(
                modality="document",
                summary=f"Document not found: {document_path}",
                facts=[],
                evidence=[],
                confidence=0.0,
                missing_info=["Please check document path"],
                raw_output="",
            )

        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".csv", ".tsv"}:
            text = self._read_txt(path)
        elif suffix == ".pdf":
            text = self._read_pdf(path)
        elif suffix == ".docx":
            text = self._read_docx(path)
        elif suffix in {".xlsx", ".xls"}:
            text = self._read_xlsx(path)
        else:
            return PerceptionResult(
                modality="document",
                summary=f"Unsupported document format: {suffix}",
                facts=[],
                evidence=[],
                confidence=0.0,
                missing_info=[f"Document format {suffix} is not supported"],
                raw_output="",
            )

        facts = _extract_facts_fallback(text, limit=8)
        return PerceptionResult(
            modality="document",
            summary=f"Parsed {path.name}",
            facts=facts,
            evidence=[EvidenceItem(source=document_path, quote=text[:300])],
            confidence=_safe_confidence(0.6 if facts else 0.4),
            missing_info=[] if facts else ["Document has insufficient extractable information"],
            raw_output=text[:3000],
        )
