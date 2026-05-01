from __future__ import annotations

import time
import uuid
from pathlib import Path
from typing import Any, Dict, Generator, List, Optional

try:
    from .brain_client import DeepSeekBrainClient
    from .career_knowledge import CareerKnowledgeBase
    from .input_router import DataRouter
except ImportError:
    from project.core.brain_client import DeepSeekBrainClient
    from project.core.career_knowledge import CareerKnowledgeBase
    from project.core.input_router import DataRouter


class PipelineError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


class MultimodalChatPipeline:
    """
    Standard conversation pipeline:
    input -> routing -> modality understanding -> RAG retrieval -> deepseek llm(stream)
    """

    def __init__(
        self,
        image_model_path: str = "./models/Qwen3-VL-2B-Instruct",
        router: Optional[DataRouter] = None,
        brain_client: Optional[DeepSeekBrainClient] = None,
        image_processor: Any = None,
        image_agent: Any = None,
        doc_agent: Any = None,
        audio_agent: Any = None,
        video_agent: Any = None,
        knowledge_base: Any = None,
    ) -> None:
        self.router = router or DataRouter()
        self.brain_client = brain_client or DeepSeekBrainClient()
        self.knowledge = knowledge_base or CareerKnowledgeBase()
        self.image_model_path = image_model_path
        self._image_processor = image_processor
        self._image_agent = image_agent
        self._doc_agent = doc_agent
        self._audio_agent = audio_agent
        self._video_agent = video_agent
        self._session_history: Dict[str, List[Dict[str, str]]] = {}

    def _get_image_agent(self):
        if self._image_agent is None:
            try:
                from project.agents.perception import ImagePerceptionAgent
            except Exception:
                from agents.perception import ImagePerceptionAgent
            self._image_agent = ImagePerceptionAgent(model_path=self.image_model_path)
        return self._image_agent

    def _get_doc_agent(self):
        if self._doc_agent is not None:
            return self._doc_agent
        try:
            from project.agents.perception import DocumentPerceptionAgent
        except Exception:
            from agents.perception import DocumentPerceptionAgent
        self._doc_agent = DocumentPerceptionAgent()
        return self._doc_agent

    @staticmethod
    def _event(
        event: str,
        trace_id: str,
        session_id: str,
        stage: str,
        stage_start: float,
        data: Dict[str, Any],
    ) -> Dict[str, Any]:
        payload = dict(data)
        payload["trace_id"] = trace_id
        payload["session_id"] = session_id
        payload["stage"] = stage
        payload["elapsed_ms"] = int((time.perf_counter() - stage_start) * 1000)
        return {"event": event, "data": payload}

    def _understand_images(self, image_paths: List[str], user_text: str = "") -> str:
        agent = self._get_image_agent()
        chunks: List[str] = []
        for path in image_paths:
            result = agent.perceive(path, user_goal=user_text, user_text=user_text)
            chunks.append(f"[{path}] {result.raw_output}")
        agent.unload()
        return "\n".join(chunks).strip()

    @staticmethod
    def _read_text_file(path: Path) -> str:
        return path.read_text(encoding="utf-8", errors="ignore")

    @staticmethod
    def _read_pdf_file(path: Path) -> str:
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception:
            return f"[{path}] PDF parser unavailable (install pypdf)."
        reader = PdfReader(str(path))
        parts: List[str] = []
        for page in reader.pages[:10]:
            parts.append(page.extract_text() or "")
        text = "\n".join(parts).strip()
        return text if text else f"[{path}] PDF has no extractable text."

    @staticmethod
    def _read_docx_file(path: Path) -> str:
        try:
            import docx  # type: ignore
        except Exception:
            return f"[{path}] DOCX parser unavailable (install python-docx)."
        doc = docx.Document(str(path))
        text = "\n".join([p.text for p in doc.paragraphs]).strip()
        return text if text else f"[{path}] DOCX has no extractable text."

    @staticmethod
    def _read_xlsx_file(path: Path) -> str:
        try:
            import openpyxl  # type: ignore
        except Exception:
            return f"[{path}] XLSX parser unavailable (install openpyxl)."
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        lines: List[str] = []
        for ws in wb.worksheets[:3]:
            lines.append(f"[sheet] {ws.title}")
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                values = [str(x) for x in row if x is not None and str(x).strip()]
                if values:
                    lines.append(" | ".join(values[:12]))
                row_count += 1
                if row_count >= 20:
                    break
        wb.close()
        return "\n".join(lines).strip() if lines else f"[{path}] XLSX has no extractable content."

    def _understand_documents(self, document_paths: List[str], user_text: str = "") -> str:
        agent = self._get_doc_agent()
        chunks: List[str] = []
        for raw_path in document_paths:
            result = agent.perceive(raw_path)
            content = result.raw_output if result.raw_output else result.summary
            if len(content) > 6000:
                content = content[:6000] + "\n...[truncated]"
            path = Path(raw_path)
            chunks.append(f"[document:{path.name}]\n{content}")

        combined = "\n\n".join(chunks).strip()
        if user_text:
            return f"User context:\n{user_text}\n\nDocument extraction:\n{combined}"
        return combined

    def _get_audio_agent(self):
        if self._audio_agent is not None:
            return self._audio_agent
        try:
            from project.agents.perception.audio_agent import AudioPerceptionAgent
        except Exception:
            from agents.perception.audio_agent import AudioPerceptionAgent
        self._audio_agent = AudioPerceptionAgent()
        return self._audio_agent

    def _get_video_agent(self):
        if self._video_agent is not None:
            return self._video_agent
        try:
            from project.agents.perception.video_agent import VideoPerceptionAgent
        except Exception:
            from agents.perception.video_agent import VideoPerceptionAgent
        self._video_agent = VideoPerceptionAgent(
            image_model_path=self.image_model_path
        )
        return self._video_agent

    def _understand_audio_video(self, media_paths: List[str], user_text: str = "") -> str:
        audio_agent = self._get_audio_agent()
        video_agent = None
        chunks: List[str] = []
        for raw_path in media_paths:
            path = Path(raw_path)
            suffix = path.suffix.lower()
            if suffix in {".mp3", ".wav", ".flac", ".aac", ".ogg", ".m4a"}:
                try:
                    perceived = audio_agent.perceive(raw_path)
                    info = perceived.raw_output.strip() or perceived.summary
                    if len(info) > 3000:
                        info = info[:3000] + "\n...[truncated]"
                    chunks.append(f"[audio:{path.name}] {info}")
                except Exception as exc:
                    chunks.append(f"[audio:{path.name}] audio processing failed: {exc}")
            else:
                if video_agent is None:
                    video_agent = self._get_video_agent()
                try:
                    perceived = video_agent.perceive(raw_path)
                    info = perceived.raw_output.strip() or perceived.summary
                    if len(info) > 3000:
                        info = info[:3000] + "\n...[truncated]"
                    chunks.append(f"[video:{path.name}] {info}")
                except Exception as exc:
                    chunks.append(f"[video:{path.name}] video processing failed: {exc}")

        combined = "\n".join(chunks).strip()
        if user_text:
            return f"User context:\n{user_text}\n\nAudio/Video extraction:\n{combined}"
        return combined

    def _small_model_understanding(self, routed: Dict[str, Any]) -> str:
        mode = routed["route"]
        payload = routed["payload"]

        # As requested for the final target: text should go to LLM directly.
        if mode == "text":
            return payload.get("text", "")
        if mode == "image":
            return self._understand_images(payload.get("images", []), "")
        if mode == "multimodal":
            return self._understand_images(payload.get("images", []), payload.get("text", ""))
        if mode == "file":
            return self._understand_documents(payload.get("documents", []), payload.get("text", ""))
        if mode == "audio_video":
            return self._understand_audio_video(payload.get("media", []), payload.get("text", ""))
        raise PipelineError("UNKNOWN_MODE", f"unsupported mode: {mode}")

    def _history_as_text(self, session_id: str, max_turns: int = 6) -> str:
        history = self._session_history.get(session_id, [])
        if not history:
            return ""
        lines: List[str] = []
        for turn in history[-max_turns:]:
            lines.append(f"user: {turn.get('user', '')}")
            lines.append(f"assistant: {turn.get('assistant', '')}")
        return "\n".join(lines).strip()

    def _build_llm_prompt(
        self,
        user_input: str,
        routed: Dict[str, Any],
        understanding: str,
        session_id: str,
        references: List[Dict[str, str]],
    ) -> str:
        history_text = self._history_as_text(session_id)
        history_block = f"Recent conversation history:\n{history_text}\n\n" if history_text else ""
        references_text = "\n".join(
            [
                (
                    f"- role: {x.get('role', '')}; skills: {x.get('skills', '')}; "
                    f"resources: {x.get('resources', '')}; paths: {x.get('transition_paths', '')}; "
                    f"salary: {x.get('salary_hint', '')}; score: {x.get('match_score', '')}"
                )
                for x in references
            ]
        )
        return (
            "You are a professional career planning assistant. "
            "Use the retrieved references as grounding evidence.\n\n"
            f"{history_block}"
            "Pipeline context:\n"
            f"- route_mode: {routed['route']}\n"
            f"- target_agent: {routed['target_agent']}\n"
            f"- modality_understanding:\n{understanding}\n\n"
            "Retrieved knowledge references:\n"
            f"{references_text}\n\n"
            "User question:\n"
            f"{user_input}\n"
        )

    def get_session_history(self, session_id: str) -> List[Dict[str, str]]:
        return list(self._session_history.get(session_id, []))

    def clear_session(self, session_id: str) -> None:
        self._session_history.pop(session_id, None)

    def run_stream(
        self,
        user_input: str,
        llm_model: Optional[str] = None,
        session_id: str = "default",
        trace_id: Optional[str] = None,
    ) -> Generator[Dict[str, Any], None, None]:
        trace_id = trace_id or str(uuid.uuid4())

        routing_t0 = time.perf_counter()
        routed = self.router.route(user_input)
        if routed.get("error"):
            raise PipelineError("ROUTING_FAILED", routed["error"])

        mode = routed["route"]
        valid, msg = self.router.classifier.validate(routed["metadata"])
        if not valid:
            raise PipelineError("VALIDATION_FAILED", msg)

        yield self._event(
            "route",
            trace_id,
            session_id,
            "routing",
            routing_t0,
            {"mode": mode, "target_agent": routed["target_agent"]},
        )

        understanding_t0 = time.perf_counter()
        understanding = self._small_model_understanding(routed)
        yield self._event(
            "small_model_output",
            trace_id,
            session_id,
            "small_model",
            understanding_t0,
            {"mode": mode, "text": understanding},
        )

        rag_t0 = time.perf_counter()
        references = self.knowledge.retrieve(f"{user_input}\n{understanding}", top_k=4)
        yield self._event(
            "rag_references",
            trace_id,
            session_id,
            "rag",
            rag_t0,
            {"count": len(references), "items": references},
        )

        prompt_t0 = time.perf_counter()
        llm_prompt = self._build_llm_prompt(
            user_input=user_input,
            routed=routed,
            understanding=understanding,
            session_id=session_id,
            references=references,
        )
        yield self._event(
            "llm_input_ready",
            trace_id,
            session_id,
            "llm_prompt",
            prompt_t0,
            {"prompt_preview": llm_prompt[:1200]},
        )

        llm_t0 = time.perf_counter()
        full_text = ""
        for token in self.brain_client.plan_stream(llm_prompt, model=llm_model):
            full_text += token
            yield self._event(
                "token",
                trace_id,
                session_id,
                "llm_stream",
                llm_t0,
                {"token": token},
            )

        self._session_history.setdefault(session_id, []).append({"user": user_input, "assistant": full_text})

        yield self._event(
            "final",
            trace_id,
            session_id,
            "finalize",
            llm_t0,
            {"text": full_text, "mode": mode, "references_used": len(references)},
        )
